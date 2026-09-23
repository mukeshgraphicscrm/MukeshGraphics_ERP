import openpyxl
import sys

def parse_excel(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=False)
    sheet = wb.active
    
    print("## Fixed Values")
    for row in range(13, 22):
        for col in range(2, 7):
            cell = sheet.cell(row=row, column=col)
            if cell.value is not None:
                print(f"{cell.coordinate}: {cell.value}")
                
    print("\n## Main Calculation Table Headers")
    for col in range(2, 10):
        cell = sheet.cell(row=21, column=col)
        print(f"{cell.coordinate}: {cell.value}")
        
    print("\n## Main Calculation Table First Row (Row 22)")
    for col in range(2, 10):
        cell = sheet.cell(row=22, column=col)
        print(f"{cell.coordinate}: {cell.value}")
        
    print("\n## Lid Size Data Headers (L to AZ)")
    # Just print the first one or two groups to see the pattern
    print("Group 1: 72mm")
    for row in range(18, 23):
        for col in range(12, 16):
            cell = sheet.cell(row=row, column=col)
            if cell.value is not None:
                print(f"{cell.coordinate}: {cell.value}")

if __name__ == "__main__":
    parse_excel("C:\\Internship DZ Infotech\\Mukesh Graphics Erp - Copy\\LID RATE.xlsx")
