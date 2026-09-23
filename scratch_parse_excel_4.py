import openpyxl

def dump_headers(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=False)
    sheet = wb.active
    
    print("## Lid Sizes")
    for col in range(12, 53, 4): # L, P, T, X, AB, AF, AJ, AN, AR, AV, AZ... step by 4
        size_cell = sheet.cell(row=18, column=col)
        factor_cell = sheet.cell(row=20, column=col)
        
        if size_cell.value:
            print(f"Size: {size_cell.value}, Factor: {factor_cell.value}")

if __name__ == "__main__":
    dump_headers("C:\\Internship DZ Infotech\\Mukesh Graphics Erp - Copy\\LID RATE.xlsx")
